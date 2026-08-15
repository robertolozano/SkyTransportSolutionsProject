"""
Generates the submission deck.

Kept in the repo so the deck is reproducible and reviewable rather than a binary
someone has to trust. Run:  /tmp/deckenv/bin/python scripts/build-deck.py
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# Palette lifted from the application's design tokens, so the deck and the
# product read as one thing.
NAVY = RGBColor(0x1E, 0x3A, 0x8A)
INK = RGBColor(0x0F, 0x17, 0x2A)
INK_SOFT = RGBColor(0x47, 0x55, 0x69)
FAINT = RGBColor(0x94, 0xA3, 0xB8)
CANVAS = RGBColor(0xF6, 0xF7, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CRITICAL = RGBColor(0xB9, 0x1C, 0x1C)
GOOD = RGBColor(0x15, 0x80, 0x3D)
HIGH = RGBColor(0xC2, 0x41, 0x0C)
EDGE = RGBColor(0xE2, 0xE8, 0xF0)

FONT = "Arial"
W, H = Inches(13.333), Inches(7.5)

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]


def bg(slide, color=WHITE):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def box(slide, left, top, width, height, fill=None, line=None):
    from pptx.enum.shapes import MSO_SHAPE
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shp.adjustments[0] = 0.06
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def text(slide, s, left, top, width, height, size=18, color=INK, bold=False,
         align=PP_ALIGN.LEFT, space_after=6, line_spacing=1.15):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    lines = s if isinstance(s, list) else [s]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        p.line_spacing = line_spacing
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = FONT
    return tb


def bullets(slide, items, left, top, width, size=16, color=INK_SOFT, gap=10):
    """items: list of (bold_lead, rest) or plain strings."""
    tb = slide.shapes.add_textbox(left, top, width, Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        p.line_spacing = 1.2
        if isinstance(item, tuple):
            lead, rest = item
            r1 = p.add_run(); r1.text = lead
            r1.font.size = Pt(size); r1.font.bold = True
            r1.font.color.rgb = INK; r1.font.name = FONT
            r2 = p.add_run(); r2.text = rest
            r2.font.size = Pt(size); r2.font.color.rgb = color; r2.font.name = FONT
        else:
            r = p.add_run(); r.text = item
            r.font.size = Pt(size); r.font.color.rgb = color; r.font.name = FONT
    return tb


def slide_with_header(title, kicker=None):
    s = prs.slides.add_slide(BLANK)
    bg(s, WHITE)
    # Navy rule at the top for a consistent spine.
    box(s, Inches(0), Inches(0), W, Inches(0.09), fill=NAVY)
    top = Inches(0.55)
    if kicker:
        text(s, kicker.upper(), Inches(0.85), top, Inches(11.5), Inches(0.3),
             size=11, color=FAINT, bold=True)
        top = Inches(0.92)
    text(s, title, Inches(0.85), top, Inches(11.6), Inches(0.9), size=30, color=INK, bold=True)
    return s


def stat(slide, left, top, value, label, color=INK, width=Inches(2.6)):
    text(slide, value, left, top, width, Inches(0.6), size=34, color=color, bold=True)
    text(slide, label, left, top + Inches(0.58), width, Inches(0.5), size=12, color=FAINT)



def two_col(slide, rows, top, left_w=Inches(4.4), size=15, row_h=Inches(0.82)):
    """Left label, right reason. For scannable pairs."""
    y = top
    for i, (label, reason) in enumerate(rows):
        if i % 2 == 0:
            box(slide, Inches(0.85), y, Inches(11.6), row_h, fill=CANVAS, line=EDGE)
        text(slide, label, Inches(1.15), y + Inches(0.22), left_w, Inches(0.4),
             size=size, color=INK, bold=True)
        text(slide, reason, Inches(1.15) + left_w, y + Inches(0.22), Inches(11.6) - left_w - Inches(0.5),
             Inches(0.4), size=size, color=INK_SOFT)
        y += row_h + Inches(0.06)
    return y


# ---------------------------------------------------------------- 1. Title
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
box(s, Inches(0), Inches(0), Inches(0.28), H, fill=NAVY)
text(s, "SKY INNOVATION  ·  AUTOMATE TRACK", Inches(1.1), Inches(2.4),
     Inches(11), Inches(0.4), size=13, color=FAINT, bold=True)
text(s, "Compliance Radar", Inches(1.1), Inches(2.9), Inches(11), Inches(1.1),
     size=54, color=INK, bold=True)
text(s, "Deadlines nobody sends you, and the ones hiding behind them.",
     Inches(1.1), Inches(4.15), Inches(11), Inches(0.5), size=22, color=INK_SOFT)
text(s, "Roberto Lozano", Inches(1.1), Inches(5.5), Inches(6), Inches(0.4),
     size=16, color=INK, bold=True)

# ---------------------------------------------------------------- 2. Stack
s = slide_with_header("What I built it with.", "The stack")
two_col(s, [
    ("Next.js + React", "Pages query the database directly. No API layer to maintain."),
    ("PostgreSQL + Prisma", "Real relations and constraints. Migrations I can review."),
    ("Raw SQL for aggregates", "Ranking and rollups read better as SQL than as ORM calls."),
    ("A pure TypeScript rules module", "No framework, no I/O — so the logic is testable and portable."),
    ("Vitest + Playwright", "Fast tests with no database; a browser only where it's needed."),
], Inches(2.05))
text(s, "One deliberate constraint: the compliance logic imports nothing — so it tests without fixtures and runs in the browser unchanged.",
     Inches(0.85), Inches(6.65), Inches(11.6), Inches(0.5), size=14, color=INK_SOFT)

# ---------------------------------------------------------------- 3. Research
s = slide_with_header("I started with the business, not the code.", "Where I began")
bullets(s, [
    "Sky Transport Solutions isn't a trucking company.",
    "They're the outsourced paperwork department for ~18,000 carriers — about 25 people.",
    "Their customers are owner-operators. The guy who owns the truck is driving it.",
    "Memberships run $199–$399 per truck per year, tiered by how many renewals Sky remembers for you.",
], Inches(0.85), Inches(2.1), Inches(11.6), size=18, gap=16)
box(s, Inches(0.85), Inches(4.9), Inches(11.6), Inches(1.2), fill=CANVAS, line=EDGE)
text(s, "So they sell memory. That reframed the whole problem.",
     Inches(1.2), Inches(5.3), Inches(11), Inches(0.5), size=20, color=NAVY, bold=True)

# ---------------------------------------------------------------- 4. Pain points
s = slide_with_header("Where it hurts.", "What I found")
pains = [
    ("Nobody sends a bill.", "Deadlines are derived. One is encoded in the carrier's own DOT number."),
    ("Filings block each other.", "No tax receipt, no plate renewal. No emissions check, no registration."),
    ("The client is driving.", "You can't collect a document from someone on I-5."),
    ("Everything lands at once.", "Four fuel tax dates a year, on every client, simultaneously."),
    ("One miss costs more than the fee.", "$199 to remember. ~$800 a day when a truck is parked."),
]
y = Inches(2.05)
for label, detail in pains:
    box(s, Inches(0.85), y, Inches(11.6), Inches(0.82), fill=CANVAS, line=EDGE)
    text(s, label, Inches(1.15), y + Inches(0.22), Inches(4.3), Inches(0.4), size=15, color=INK, bold=True)
    text(s, detail, Inches(5.5), y + Inches(0.22), Inches(6.7), Inches(0.4), size=15, color=INK_SOFT)
    y += Inches(0.88)

# ---------------------------------------------------------------- 5. What I picked
s = slide_with_header("I picked three.", "Scope")
picks = [
    ("Derive the deadlines", "Encode the regulations, don't store dates.", NAVY),
    ("Model what blocks what", "The chain is the part no calendar shows.", HIGH),
    ("Collect the documents", "Ask early, chase automatically, read what arrives.", GOOD),
]
x = Inches(0.85)
for title, sub, accent in picks:
    box(s, x, Inches(2.2), Inches(3.65), Inches(2.5), fill=CANVAS, line=EDGE)
    box(s, x, Inches(2.2), Inches(3.65), Inches(0.08), fill=accent)
    text(s, title, x + Inches(0.3), Inches(2.6), Inches(3.05), Inches(0.5), size=19, color=INK, bold=True)
    text(s, sub, x + Inches(0.3), Inches(3.25), Inches(3.05), Inches(1.1), size=14, color=INK_SOFT)
    x += Inches(3.9)
text(s, "The rest — seasonal planning, filing packets, coverage gaps — fell out of the same engine once it existed.",
     Inches(0.85), Inches(5.2), Inches(11.6), Inches(0.5), size=15, color=INK_SOFT)

# ---------------------------------------------------------------- 6. Deriving
s = slide_with_header("Deriving the deadlines.", "How I addressed it — 1 of 3")
bullets(s, [
    "Seven rules, each carrying the regulation it came from.",
    "One is derived from digits in the DOT number. Nobody is ever told that date.",
    "Two are marked \"needs review\" — I wasn't certain, so the app says so.",
    "Ambiguous dates are refused, not guessed. 03/04 could be March or April.",
], Inches(0.85), Inches(2.1), Inches(11.6), size=18, gap=16)
box(s, Inches(0.85), Inches(4.9), Inches(11.6), Inches(1.3), fill=CANVAS, line=EDGE)
text(s, "A rule that looks confident and is wrong is worse than one that admits it isn't sure.",
     Inches(1.2), Inches(5.35), Inches(11), Inches(0.5), size=17, color=NAVY, bold=True)

# ---------------------------------------------------------------- 7. The chain
s = slide_with_header("The date nobody's calendar contains.", "How I addressed it — 2 of 3")
chain = [("Emissions check", "Aug 17", CRITICAL), ("Heavy vehicle tax", "Aug 31", HIGH),
         ("Plate renewal", "Sep 16", INK)]
x = Inches(0.85)
for i, (label, date, accent) in enumerate(chain):
    box(s, x, Inches(2.25), Inches(3.1), Inches(1.25), fill=CANVAS, line=EDGE)
    box(s, x, Inches(2.25), Inches(3.1), Inches(0.07), fill=accent)
    text(s, label, x + Inches(0.25), Inches(2.55), Inches(2.6), Inches(0.4), size=15, color=INK, bold=True)
    text(s, date, x + Inches(0.25), Inches(2.95), Inches(2.6), Inches(0.4), size=14, color=INK_SOFT)
    if i < 2:
        text(s, "→", x + Inches(3.2), Inches(2.7), Inches(0.5), Inches(0.4), size=20, color=FAINT)
    x += Inches(3.7)
text(s, "TRUCK STOPS BEING LEGAL", Inches(8.25), Inches(3.62), Inches(3.5), Inches(0.3),
     size=11, color=CRITICAL, bold=True)
bullets(s, [
    "The plate looks like a September problem.",
    "It's an August one — miss the emissions check and the renewal can't complete.",
    "That truck is on the plan that doesn't cover plate renewals.",
], Inches(0.85), Inches(4.35), Inches(11.6), size=18, gap=14)

# ---------------------------------------------------------------- 8. Documents
s = slide_with_header("Getting documents off a driver's phone.", "How I addressed it — 3 of 3")
bullets(s, [
    "We ask 60 days before expiry, then chase on a fixed ladder. Texts first — he's driving.",
    "He photographs it at a truck stop. No app, no login.",
    "The image is shrunk in the browser, because signal out there is bad.",
    "A model reads it. Then a person confirms it before anything is filed.",
], Inches(0.85), Inches(2.1), Inches(11.6), size=18, gap=16)
box(s, Inches(0.85), Inches(4.95), Inches(11.6), Inches(1.3), fill=CANVAS, line=EDGE)
text(s, "Extraction proposes. A person commits. A misread expiry date is the exact failure this thing exists to prevent.",
     Inches(1.2), Inches(5.4), Inches(11), Inches(0.5), size=17, color=NAVY, bold=True)

# ---------------------------------------------------------------- 9. AI workflow
s = slide_with_header("Research first, then encode it.", "How I worked with AI")
bullets(s, [
    "I learned the domain before writing anything. That's why the rules carry citations.",
    "Then I turned that context into five subagents, committed to the repo.",
    "Compliance, data layer, security, testing, design — each holds judgment the code doesn't.",
    "I rejected five more. \"Look for bugs\" isn't knowledge, it's default behaviour.",
], Inches(0.85), Inches(2.1), Inches(11.6), size=18, gap=16)
box(s, Inches(0.85), Inches(4.95), Inches(11.6), Inches(1.3), fill=CANVAS, line=EDGE)
text(s, "Deciding what not to build was the harder half.",
     Inches(1.2), Inches(5.4), Inches(11), Inches(0.5), size=18, color=NAVY, bold=True)

# ---------------------------------------------------------------- 10. Testing
s = slide_with_header("The tests caught two real bugs.", "What happened")
box(s, Inches(0.85), Inches(2.1), Inches(5.65), Inches(3.1), fill=CANVAS, line=EDGE)
box(s, Inches(0.85), Inches(2.1), Inches(5.65), Inches(0.08), fill=CRITICAL)
text(s, "Valid PDFs were rejected", Inches(1.15), Inches(2.45), Inches(5.05), Inches(0.4),
     size=17, color=INK, bold=True)
text(s, "I trusted the browser's file type. It's guessed from the extension, so a file without one reported nothing.\n\nNow the server reads the first few bytes instead.",
     Inches(1.15), Inches(3.0), Inches(5.05), Inches(2.0), size=14, color=INK_SOFT)
box(s, Inches(6.8), Inches(2.1), Inches(5.65), Inches(3.1), fill=CANVAS, line=EDGE)
box(s, Inches(6.8), Inches(2.1), Inches(5.65), Inches(0.08), fill=CRITICAL)
text(s, "The simulator went quiet", Inches(7.1), Inches(2.45), Inches(5.05), Inches(0.4),
     size=17, color=INK, bold=True)
text(s, "Slipping a deadline past the one it blocks deleted the link instead of breaking it — so it reported no problem.\n\nIn exactly the case it exists to show.",
     Inches(7.1), Inches(3.0), Inches(5.05), Inches(2.0), size=14, color=INK_SOFT)
text(s, "111 tests. The interesting ones pin decisions — like refusing an ambiguous date — not implementations.",
     Inches(0.85), Inches(5.5), Inches(11.6), Inches(0.5), size=15, color=INK_SOFT)

# ---------------------------------------------------------------- 11. Results
s = slide_with_header("Where it landed.", "Results")
stat(s, Inches(0.85), Inches(2.2), "7", "rules, each cited")
stat(s, Inches(3.45), Inches(2.2), "22", "routes")
stat(s, Inches(6.05), Inches(2.2), "111", "tests")
stat(s, Inches(8.65), Inches(2.2), "1,150", "deadlines derived")
bullets(s, [
    "Three surfaces: a public walkthrough, a client portal, and the staff console.",
    "Live on Vercel and Neon.",
], Inches(0.85), Inches(3.9), Inches(11.6), size=18, gap=14)
box(s, Inches(0.85), Inches(5.0), Inches(11.6), Inches(1.35), fill=CANVAS, line=EDGE)
text(s, "Next: a live FMCSA lookup, so it runs on a real carrier instead of seeded data.",
     Inches(1.2), Inches(5.3), Inches(11), Inches(0.45), size=16, color=NAVY, bold=True)
text(s, "And letting a confirmed document update the record, which today stops at the comparison.",
     Inches(1.2), Inches(5.78), Inches(11), Inches(0.45), size=14, color=INK_SOFT)

# ---------------------------------------------------------------- 12. Demo
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
box(s, Inches(0), Inches(0), Inches(0.28), H, fill=NAVY)
text(s, "Let me show you.", Inches(1.1), Inches(2.6), Inches(11), Inches(0.9),
     size=46, color=INK, bold=True)
steps = [
    "The walkthrough — answer one question differently, watch the requirements halve",
    "The work queue — and what a parked truck actually costs",
    "One truck — the chain, and the slider that breaks it",
    "The client's phone — send a document, watch it get read back",
]
top = Inches(3.9)
for i, stp in enumerate(steps):
    text(s, str(i + 1), Inches(1.1), top, Inches(0.4), Inches(0.4), size=15, color=NAVY, bold=True)
    text(s, stp, Inches(1.6), top, Inches(10.6), Inches(0.4), size=17, color=INK_SOFT)
    top += Inches(0.6)

import os
out = "docs/compliance-radar-deck.pptx"
os.makedirs("docs", exist_ok=True)
prs.save(out)
print(f"wrote {out} — {len(prs.slides._sldIdLst)} slides")
