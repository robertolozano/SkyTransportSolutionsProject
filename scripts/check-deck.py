"""
Layout check for the generated deck.

python-pptx will happily write text that runs off the slide, and the failure is
invisible until someone opens it in front of an interviewer. This estimates the
rendered height of every text frame and flags anything that overflows its slide
or collides with the shape below it.

Estimates, not exact metrics — but gross overflow is what actually happens, and
that it catches reliably.
"""
import math
import sys
from pptx import Presentation
from pptx.util import Emu

DECK = "docs/compliance-radar-deck.pptx"

# Arial: average advance width is ~0.5 em for mixed-case prose; bold a little wider.
AVG_CHAR_EM = 0.47
BOLD_CHAR_EM = 0.51
LINE_HEIGHT = 1.20

prs = Presentation(DECK)
slide_h_pt = prs.slide_height / 12700
slide_w_pt = prs.slide_width / 12700

problems = []
checked = 0

for index, slide in enumerate(prs.slides, start=1):
    for shape in slide.shapes:
        if not shape.has_text_frame:
            # Still confirm the shape itself is on the slide.
            right = (shape.left + shape.width) / 12700
            bottom = (shape.top + shape.height) / 12700
            if bottom > slide_h_pt + 1 or right > slide_w_pt + 1:
                problems.append(f"slide {index}: shape extends off-slide (bottom {bottom:.0f}pt)")
            continue

        tf = shape.text_frame
        width_pt = shape.width / 12700
        top_pt = shape.top / 12700
        if width_pt <= 0:
            continue

        total_h = 0.0
        for para in tf.paragraphs:
            runs = para.runs
            if not runs:
                total_h += 12 * LINE_HEIGHT
                continue
            size = max((r.font.size.pt if r.font.size else 18) for r in runs)
            bold = any(r.font.bold for r in runs)
            em = BOLD_CHAR_EM if bold else AVG_CHAR_EM
            chars = sum(len(r.text) for r in runs)
            # Explicit newlines force a break regardless of width.
            hard_lines = sum(r.text.count("\n") for r in runs)
            per_line = max(1, int(width_pt / (em * size)))
            lines = max(1, math.ceil(chars / per_line)) + hard_lines
            spacing = para.line_spacing if isinstance(para.line_spacing, float) else 1.0
            total_h += lines * size * LINE_HEIGHT * spacing
            total_h += para.space_after.pt if para.space_after else 0

        checked += 1
        bottom = top_pt + total_h
        if bottom > slide_h_pt - 8:  # keep an 8pt bottom margin
            preview = (tf.text[:52] + "…") if len(tf.text) > 52 else tf.text
            problems.append(
                f"slide {index}: text overflows bottom by {bottom - slide_h_pt:+.0f}pt "
                f"(ends {bottom:.0f}pt of {slide_h_pt:.0f}pt) — \"{preview.strip()}\""
            )

# Collision detection. Overflow past the slide edge is only half the failure —
# a bold heading that wraps to two lines and lands on the paragraph beneath it
# stays inside the slide and still looks broken. That is a real bug this deck
# had, and the reason this second pass exists.
for index, slide in enumerate(prs.slides, start=1):
    frames = []
    for shape in slide.shapes:
        if not shape.has_text_frame or not shape.text_frame.text.strip():
            continue
        width_pt = shape.width / 12700
        if width_pt <= 0:
            continue
        h = 0.0
        for para in shape.text_frame.paragraphs:
            runs = para.runs
            if not runs:
                continue
            size = max((r.font.size.pt if r.font.size else 18) for r in runs)
            bold = any(r.font.bold for r in runs)
            em = BOLD_CHAR_EM if bold else AVG_CHAR_EM
            chars = sum(len(r.text) for r in runs)
            per_line = max(1, int(width_pt / (em * size)))
            lines = max(1, math.ceil(chars / per_line)) + sum(r.text.count("\n") for r in runs)
            spacing = para.line_spacing if isinstance(para.line_spacing, float) else 1.0
            h += lines * size * LINE_HEIGHT * spacing
        frames.append((
            shape.top / 12700, shape.top / 12700 + h,
            shape.left / 12700, shape.left / 12700 + width_pt,
            shape.text_frame.text[:38].replace("\n", " ").strip(),
        ))

    for i in range(len(frames)):
        for j in range(i + 1, len(frames)):
            t1, b1, l1, r1, x1 = frames[i]
            t2, b2, l2, r2, x2 = frames[j]
            horizontal = l1 < r2 - 2 and l2 < r1 - 2
            vertical = t1 < b2 - 7 and t2 < b1 - 7
            if horizontal and vertical:
                problems.append(
                    f"slide {index}: text collides — \"{x1}\" overlaps \"{x2}\""
                )

print(f"{len(prs.slides._sldIdLst)} slides, {checked} text frames checked")
if problems:
    print(f"\n{len(problems)} layout problem(s):\n")
    for p in problems:
        print("  " + p)
    sys.exit(1)
print("no overflow detected")
